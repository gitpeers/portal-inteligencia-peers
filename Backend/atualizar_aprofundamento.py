"""Monta a visão por player dos Indicadores setoriais: para cada indicador ligado em config/indicadores_players.json,
as empresas (ou os entes públicos) que se comparam a ele, com a métrica equivalente e a diferença para o setor.

Sem IA. Lê os dados que os outros scripts já gravaram (indicadores.json, companhias_abertas.json, valor1000.json) e
grava aprofundamento.json; a rotina de IA escreve a leitura de cada indicador por cima (Backend/ia/ROTINA.md).
Quando o indicador tem série mensal, a comparação usa os mesmos meses do trimestre das empresas.

    python Backend/atualizar_aprofundamento.py
"""
import csv
import io
from collections import Counter
from datetime import date

from comum import CONFIG, DADOS, agora, baixar, gravar_json, ler_json, log, registrar_atualizacao

EM_LINHA = 0.5            # diferença menor que isso (em p.p.) conta como "em linha com o setor"
MAX_PADRAO = 10
CAPAG_VALIDADE_DIAS = 30  # as notas Capag mudam poucas vezes por ano: baixa de novo depois disso

METRICAS = {   # rótulo e unidade de cada métrica das empresas
    "cresc": ("Crescimento da receita", "%"),
    "crescCarteira": ("Crescimento da carteira de crédito", "%"),
    "inadimplencia": ("Inadimplência acima de 90 dias", "%"),
    "crescLL": ("Crescimento do lucro", "%"),
    "cet1": ("Capital principal (CET1)", "%"),
    "lucroLiquido": ("Lucro líquido", "R$ mi"),
    "capag": ("Nota Capag", ""),
}
CAPITAIS = {   # código IBGE das capitais (Brasília não tem Capag municipal)
    1100205: "Porto Velho", 1200401: "Rio Branco", 1302603: "Manaus", 1400100: "Boa Vista", 1501402: "Belém",
    1600303: "Macapá", 1721000: "Palmas", 2111300: "São Luís", 2211001: "Teresina", 2304400: "Fortaleza",
    2408102: "Natal", 2507507: "João Pessoa", 2611606: "Recife", 2704302: "Maceió", 2800308: "Aracaju",
    2927408: "Salvador", 3106200: "Belo Horizonte", 3205309: "Vitória", 3304557: "Rio de Janeiro", 3550308: "São Paulo",
    4106902: "Curitiba", 4205407: "Florianópolis", 4314902: "Porto Alegre", 5002704: "Campo Grande",
    5103403: "Cuiabá", 5208707: "Goiânia",
}
UFS = {"AC": "Acre", "AL": "Alagoas", "AP": "Amapá", "AM": "Amazonas", "BA": "Bahia", "CE": "Ceará", "DF": "Distrito Federal",
       "ES": "Espírito Santo", "GO": "Goiás", "MA": "Maranhão", "MT": "Mato Grosso", "MS": "Mato Grosso do Sul",
       "MG": "Minas Gerais", "PA": "Pará", "PB": "Paraíba", "PR": "Paraná", "PE": "Pernambuco", "PI": "Piauí",
       "RJ": "Rio de Janeiro", "RN": "Rio Grande do Norte", "RS": "Rio Grande do Sul", "RO": "Rondônia", "RR": "Roraima",
       "SC": "Santa Catarina", "SP": "São Paulo", "SE": "Sergipe", "TO": "Tocantins"}
ORDEM_CAPAG = {"A+": 0, "A": 1, "B+": 2, "B": 3, "C": 4, "D": 5}
TRIMESTRE_MESES = {"1": (1, 2, 3), "2": (4, 5, 6), "3": (7, 8, 9), "4": (10, 11, 12)}


# ---------- empresas ----------

def empresas_abertas(valor_setor):
    for e in ler_json(DADOS / "companhias_abertas.json")["empresas"]:
        m = e.get("metr", {})
        yield {"nome": e["nome"], "div": e["div"], "setorValor": valor_setor.get(e.get("valorId") or ""),
               "tamanho": m.get("receita") or m.get("carteira") or 0, "metr": m, "periodo": e.get("periodo", ""),
               "fonte": "Release de resultados" if e.get("fonte") == "release" else "CVM (ITR)", "temValor": bool(e.get("valorId"))}


def empresas_valor(linhas, campos):
    for linha in linhas:
        v = dict(zip(campos, linha))
        yield {"nome": v["nome"], "div": v["div"], "setorValor": v["setor"], "tamanho": v.get("rol") or 0,
               "metr": {"cresc": v.get("varRol"), "lucroLiquido": v.get("ll")}, "periodo": "2025",
               "fonte": "Valor 1000 (2025)", "temValor": True}


def escolher(universo, filtro):
    """Empresas do filtro: a divisão e o setor do Valor delimitam; 'nomes' inclui pelo nome; 'excluir' tira."""
    nomes, excluir = set(filtro.get("nomes", [])), set(filtro.get("excluir", []))
    escolhidas = []
    for e in universo:
        if e["nome"] in excluir:
            continue
        por_nome = e["nome"] in nomes
        por_filtro = ("div" in filtro or "setorValor" in filtro) and e["temValor"] \
            and filtro.get("div", e["div"]) == e["div"] \
            and ("setorValor" not in filtro or e["setorValor"] in filtro["setorValor"])
        if por_nome or por_filtro:
            escolhidas.append(e)
    escolhidas.sort(key=lambda e: -(e["tamanho"] or 0))
    return escolhidas[:filtro.get("max", MAX_PADRAO)]


# ---------- setor no mesmo período das empresas ----------

def mes(ano, m):
    return f"{ano}-{m:02d}"


def setor_no_trimestre(ind, trimestre, agregar, comparar):
    """Valor ou variação do indicador nos mesmos meses do trimestre das empresas (ex.: 2T26 = abr a jun/26)."""
    h = ind.get("historico")
    if not (h and agregar and trimestre and len(trimestre) == 4 and trimestre[1] == "T"):
        return None
    pontos = dict(h["pontos"])
    ano = 2000 + int(trimestre[2:])
    meses = TRIMESTRE_MESES[trimestre[0]]
    if h["freq"] == "trimestral":
        meses, agregar = (meses[-1],), "fim"
    atual = [pontos.get(mes(ano, m)) for m in meses]
    antes = [pontos.get(mes(ano - 1, m)) for m in meses]
    if comparar == "valor":
        return atual[-1]
    if None in atual or None in antes:
        return None
    if agregar == "fim":
        a, b = atual[-1], antes[-1]
    elif agregar == "soma":
        a, b = sum(atual), sum(antes)
    else:
        a, b = sum(atual) / len(atual), sum(antes) / len(antes)
    return round((a / b - 1) * 100, 1) if b else None


# ---------- entes públicos (Capag) ----------

def capag(cache):
    """Notas Capag dos estados e das capitais, guardadas por CAPAG_VALIDADE_DIAS."""
    hoje = agora().date()
    if cache.get("baixadoEm") and (hoje - date.fromisoformat(cache["baixadoEm"])).days < CAPAG_VALIDADE_DIAS:
        return cache
    from atualizar_indicadores import recurso_tt   # só aqui: o módulo é grande e só a Capag precisa dele
    import openpyxl
    novo = {"baixadoEm": hoje.isoformat()}
    url = recurso_tt("capag-estados", "capagdosestados", ultimo=False)[-1]
    linhas = list(csv.reader(io.StringIO(baixar(url).decode("latin-1")), delimiter=";"))
    i = next(k for k, c in enumerate(linhas[0]) if "CAPAG" in c.upper())
    novo["estados"] = {"ano": url.rsplit("capagdosestados", 1)[-1][:4],
                       "notas": [{"nome": f'{UFS.get(l[0].strip(), l[0].strip())} ({l[0].strip()})', "nota": l[i].strip().upper(), "parciais": [l[2].strip(), l[4].strip(), l[6].strip()]}
                                 for l in linhas[1:] if len(l) > i and l[0].strip()]}
    url = recurso_tt("capag-municipios", "capag-municipios", ultimo=False)[-1]
    wb = openpyxl.load_workbook(io.BytesIO(baixar(url, timeout=180)), read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    cab = next(k for k, l in enumerate(rows) if l and "Nome_Município" in [str(c) for c in l])
    col = {str(c): k for k, c in enumerate(rows[cab]) if c and str(c) not in ("None",)}
    capitais = []
    for l in rows[cab + 1:]:
        if l and isinstance(l[0], int) and l[0] in CAPITAIS:
            capitais.append({"nome": f'{CAPITAIS[l[0]]} ({l[col["UF"]]})', "nota": str(l[col["CAPAG"]]).strip().upper(),
                             "parciais": [str(l[col[f"Nota {k}"]] or "").strip() for k in (1, 2, 3)]})
    novo["capitais"] = {"posicao": url.rsplit("posicao-", 1)[-1].replace(".xlsx", ""), "notas": capitais}
    return novo


def players_capag(lista, fonte):
    return [{"nome": x["nome"], "valor": x["nota"], "detalhe": " · ".join(f"{n} {p or '—'}" for n, p in zip(("Endividamento", "Poupança corrente", "Liquidez"), x["parciais"])), "periodo": fonte, "fonte": "Tesouro Nacional (Capag)"}
            for x in sorted(lista, key=lambda x: (ORDEM_CAPAG.get(x["nota"], 9), x["nome"]))]


# ---------- montagem ----------

def montar(ponte, ind, universo_abertas, universo_valor, cache_capag):
    base, metrica, comparar = ponte["players"]["base"], ponte["metrica"], ponte["comparar"]
    rotulo, unidade = METRICAS[metrica]
    saida = {"segmento": ponte["segmento"], "indicador": ponte["indicador"], "nivel": ponte["nivel"], "nota": ponte["nota"],
             "metrica": metrica, "rotuloMetrica": rotulo, "unidadeMetrica": unidade, "comparar": comparar,
             "historico": ind.get("historico")}
    if base == "capag_estados":
        c = cache_capag["estados"]
        saida["players"] = players_capag(c["notas"], f'Capag {c["ano"]}')
        saida["resumo"] = Counter(p["valor"][:1] for p in saida["players"])
        return saida
    if base == "capag_capitais":
        c = cache_capag["capitais"]
        saida["players"] = players_capag(c["notas"], f'posição {c["posicao"]}')
        saida["resumo"] = Counter(p["valor"][:1] for p in saida["players"])
        return saida

    empresas = escolher(universo_abertas if base == "abertas" else universo_valor, ponte["players"])
    empresas = [e for e in empresas if isinstance(e["metr"].get(metrica), (int, float))]
    # o trimestre que a maioria das empresas divulgou (as que estão atrasadas aparecem com o próprio período)
    periodo = Counter(e["periodo"] for e in empresas).most_common(1)[0][0] if empresas else ""
    referencia, ref_periodo = None, None
    if comparar:
        alinhado = setor_no_trimestre(ind, periodo, ponte.get("agregar"), comparar)
        if alinhado is not None:
            referencia, ref_periodo = alinhado, periodo
        else:
            referencia = ind.get("var") if comparar == "var" else ind.get("valor")
            ref_periodo = ind.get("periodo")
    saida["setor"] = {"valor": ind.get("valor"), "unidade": ind.get("unidade"), "var": ind.get("var"),
                      "varPP": ind.get("varPP", False), "periodo": ind.get("periodo"), "fonte": ind.get("fonte"),
                      "referencia": referencia, "periodoReferencia": ref_periodo,
                      "mesmoPeriodo": ref_periodo == periodo and bool(periodo)}
    players = []
    for e in empresas:
        v = e["metr"][metrica]
        item = {"nome": e["nome"], "valor": v, "periodo": e["periodo"], "fonte": e["fonte"]}
        if referencia is not None:
            dif = round(v - referencia, 1)
            item["diferenca"] = dif
            item["posicao"] = "em linha" if abs(dif) < EM_LINHA else ("acima" if dif > 0 else "abaixo")
        players.append(item)
    if referencia is not None:
        players.sort(key=lambda p: -p["valor"])
    saida["players"] = players
    return saida


def main():
    pontes = ler_json(CONFIG / "indicadores_players.json")["pontes"]
    segmentos = {s["id"]: s for s in ler_json(DADOS / "indicadores.json")}
    valor = ler_json(DADOS / "valor1000.json")
    campos = valor["campos"]
    setor_por_id = {l[campos.index("id")]: l[campos.index("setor")] for l in valor["linhas"]}
    universo_abertas = list(empresas_abertas(setor_por_id))
    universo_valor = list(empresas_valor(valor["linhas"], campos))
    caminho = DADOS / "aprofundamento.json"
    anterior = ler_json(caminho) if caminho.exists() else {}
    cache = anterior.get("capag", {})
    if any(p["players"]["base"].startswith("capag") for p in pontes):
        try:
            cache = capag(cache)
        except Exception as e:   # o Tesouro fora do ar não impede o resto; vale a última nota guardada
            log(f"  falha na Capag: {str(e)[:80]}")
    saida, falhas = [], 0
    for ponte in pontes:
        seg = segmentos.get(ponte["segmento"])
        ind = next((i for i in seg["indicadores"] if i["nome"] == ponte["indicador"]), None) if seg else None
        if not ind:
            log(f"  indicador não encontrado: {ponte['segmento']} / {ponte['indicador']}")
            falhas += 1
            continue
        try:
            item = montar(ponte, ind, universo_abertas, universo_valor, cache)
        except KeyError as e:    # Capag ainda sem cache
            log(f"  sem dados para {ponte['indicador']}: {e}")
            falhas += 1
            continue
        if not item["players"]:
            log(f"  nenhum player com dados: {ponte['segmento']} / {ponte['indicador']}")
            falhas += 1
            continue
        saida.append(item)
    gravar_json(caminho, {"geradoEm": agora().isoformat(timespec="minutes"), "capag": cache, "pontes": saida})
    registrar_atualizacao("aprofundamento")
    alinhadas = sum(1 for p in saida if p.get("setor", {}).get("mesmoPeriodo"))
    log(f"Visão por player: {len(saida)} indicadores ({alinhadas} comparados no mesmo trimestre das empresas), {falhas} sem dados")


if __name__ == "__main__":
    main()
